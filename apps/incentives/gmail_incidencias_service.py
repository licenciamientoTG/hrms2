"""
Servicio para importar incidencias/correcciones de corte desde Gmail.

Flujo:
  1. Autenticación OAuth2 con el mismo token de gmail_budget_service
  2. Búsqueda del correo de walter.quezada@totalgas.com con asunto
     "CORRECCIONES E INCIDENCIAS AL <fecha>"
  3. Descarga del adjunto .xlsx
  4. Parseo de las pestañas semanales (p.ej. "27 sep", "06 sep")
     ignorando "detalle por estacion"
  5. Retorno de registros listos para guardar en IncidenciaRegistro
"""

import io
import re
import base64
import logging
import os
from datetime import date, datetime

import openpyxl
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from googleapiclient.discovery import build

logger = logging.getLogger(__name__)

SCOPES = ['https://www.googleapis.com/auth/gmail.readonly']

_BASE_DIR    = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
_CREDS_DIR   = os.path.join(_BASE_DIR, 'credentials')
_CLIENT_FILE = os.path.join(_CREDS_DIR, 'gmail_oauth_client.json')
_TOKEN_FILE  = os.path.join(_CREDS_DIR, 'gmail_token.json')

REMITENTE = 'walter.quezada@totalgas.com'

MESES_ES = {
    'ene': 1, 'feb': 2, 'mar': 3, 'abr': 4,
    'may': 5, 'jun': 6, 'jul': 7, 'ago': 8,
    'sep': 9, 'oct': 10, 'nov': 11, 'dic': 12,
    'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4,
    'mayo': 5, 'junio': 6, 'julio': 7, 'agosto': 8,
    'septiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12,
}


# ---------------------------------------------------------------------------
# Autenticación (reutiliza el mismo token que gmail_budget_service)
# ---------------------------------------------------------------------------

def _get_service():
    creds = None
    if os.path.exists(_TOKEN_FILE):
        creds = Credentials.from_authorized_user_file(_TOKEN_FILE, SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(_CLIENT_FILE, SCOPES)
            creds = flow.run_local_server(port=0)
        with open(_TOKEN_FILE, 'w') as f:
            f.write(creds.to_json())
    return build('gmail', 'v1', credentials=creds, cache_discovery=False)


# ---------------------------------------------------------------------------
# Búsqueda del correo
# ---------------------------------------------------------------------------

def _build_query(semana: date, remitente: str = None) -> str:
    """Query para una semana específica."""
    from_addr = remitente or REMITENTE
    dia = semana.day
    mes_abrev = [k for k, v in MESES_ES.items() if v == semana.month and len(k) == 3][0]
    return (
        f'from:{from_addr} '
        f'subject:"correcciones e incidencias" '
        f'subject:"{dia} de {mes_abrev}"'
    )


def _build_query_broad(remitente: str = None) -> str:
    """Query amplio para traer los correos más recientes de incidencias."""
    from_addr = remitente or REMITENTE
    return f'from:{from_addr} subject:"correcciones e incidencias"'


# ---------------------------------------------------------------------------
# Descarga del adjunto .xlsx
# ---------------------------------------------------------------------------

def _get_attachment_xlsx(service, msg_id: str) -> bytes | None:
    """Descarga el primer adjunto .xlsx del mensaje."""
    msg = service.users().messages().get(
        userId='me', id=msg_id, format='full'
    ).execute()

    def _find(parts):
        for part in parts:
            filename = part.get('filename', '')
            mime = part.get('mimeType', '')
            is_xlsx = (
                filename.lower().endswith('.xlsx')
                or mime == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            if is_xlsx:
                body = part.get('body', {})
                att_id = body.get('attachmentId')
                if att_id:
                    att = service.users().messages().attachments().get(
                        userId='me', messageId=msg_id, id=att_id
                    ).execute()
                    return base64.urlsafe_b64decode(att.get('data', ''))
            if 'parts' in part:
                result = _find(part['parts'])
                if result:
                    return result
        return None

    return _find(msg.get('payload', {}).get('parts', []))


# ---------------------------------------------------------------------------
# Parseo del Excel
# ---------------------------------------------------------------------------

def _parse_date(valor) -> date | None:
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = str(valor).strip()
    for fmt in ('%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _es_hoja_semanal(nombre: str) -> bool:
    """True si el nombre parece una fecha semanal.
    Acepta: '27 sep', '06 sep', '24 mayo 26', '07 ene 27'
    """
    n = nombre.strip().lower()
    if n == 'detalle por estacion':
        return False
    # Con año: "24 mayo 26" o "07 ene 27"
    if re.match(r'^\d{1,2}\s+[a-záéíóú]{3,}\s+\d{2}$', n):
        return True
    # Sin año: "27 sep" o "31 mayo"
    if re.match(r'^\d{1,2}\s+[a-záéíóú]{3,}$', n):
        return True
    return False


def _fecha_de_hoja(nombre: str, año: int = None) -> date | None:
    """Extrae la fecha del nombre de la hoja.
    Si el nombre incluye año de 2 dígitos (ej. '24 mayo 26'), lo usa.
    Si no, usa el año recibido como parámetro.
    """
    partes = nombre.strip().lower().split()
    try:
        dia = int(partes[0])
        mes_num = MESES_ES.get(partes[1][:3])
        if not mes_num:
            return None
        # Año explícito de 2 dígitos en la pestaña (ej. "26" → 2026)
        if len(partes) == 3:
            año = 2000 + int(partes[2])
        if not año:
            año = date.today().year
        return date(año, mes_num, dia)
    except (ValueError, TypeError, IndexError):
        return None


def _parsear_hoja(ws) -> list[dict]:
    """Parsea una hoja de detalle semanal y retorna lista de dicts."""
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []

    # Encontrar la fila de encabezados
    header_idx = None
    col_map = {}
    for i, fila in enumerate(filas):
        celdas = [
            (str(c).strip().lower()
             .replace('ó', 'o').replace('é', 'e').replace('á', 'a')
             .replace('ú', 'u').replace('í', 'i')
             if c is not None else '')
            for c in fila
        ]
        if any(c in ('problema', 'estacion') for c in celdas):
            header_idx = i
            col_map = {nombre: j for j, nombre in enumerate(celdas) if nombre}
            break

    if header_idx is None:
        return []

    def get(fila, col):
        idx = col_map.get(col)
        if idx is None:
            return None
        return fila[idx] if idx < len(fila) else None

    registros = []
    for fila in filas[header_idx + 1:]:
        if all(c is None for c in fila):
            continue
        problema = str(get(fila, 'problema') or '').strip()
        if not problema or problema.lower() in ('total general', '(en blanco)'):
            continue
        estacion_raw = str(get(fila, 'estacion') or '').strip().upper()
        if not estacion_raw:
            continue

        try:
            turno_val = int(get(fila, 'turno') or 0) or None
        except (ValueError, TypeError):
            turno_val = None

        registros.append({
            'problema':        problema,
            'fecha':           _parse_date(get(fila, 'fecha')),
            'estacion_key':    estacion_raw,
            'turno':           turno_val,
            'tipo_incidencia': str(get(fila, 'tipo de incidencia') or '').strip(),
            'sub_detalle':     str(get(fila, 'sub-detalle') or '').strip(),
            'detalle':         str(get(fila, 'detalle') or '').strip(),
            'semana':          _parse_date(get(fila, 'semana')),
            'nombre_estacion': str(get(fila, 'nombre estacion') or '').strip(),
            'mes':             _parse_date(get(fila, 'mes')),
            'zona':            str(get(fila, 'zona') or '').strip(),
        })

    return registros


def parsear_excel_incidencias(xlsx_bytes: bytes) -> dict:
    """
    Parsea el .xlsx y retorna registros agrupados por semana.

    Returns:
        {
          'semanas': { date: [ {campo: valor, ...}, ... ] },
          'error': str | None,
        }
    """
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(xlsx_bytes), read_only=True, data_only=True
        )
    except Exception as exc:
        return {'semanas': {}, 'error': f'Error al abrir el Excel: {exc}'}

    semanas = {}
    año_actual = None  # se propaga de hoja en hoja

    for nombre_hoja in wb.sheetnames:
        if not _es_hoja_semanal(nombre_hoja):
            continue

        # Año explícito en el nombre de la pestaña (ej. "24 mayo 26") → siempre confiable
        partes = nombre_hoja.strip().lower().split()
        if len(partes) == 3:
            try:
                año_actual = 2000 + int(partes[2])
            except (ValueError, TypeError):
                pass

        ws = wb[nombre_hoja]
        registros = _parsear_hoja(ws)
        if not registros:
            continue

        # El año real lo determinamos desde la columna "Fecha" de los registros
        # (fecha del evento), que siempre trae dd/mm/yyyy completo y es la más confiable.
        # La columna "Semana" a veces tiene el año mal capturado en el Excel.
        año_real = None
        for reg in registros:
            if reg.get('fecha'):
                año_real = reg['fecha'].year
                break
        if año_real is None:
            # Fallback: columna Mes
            for reg in registros:
                if reg.get('mes'):
                    año_real = reg['mes'].year
                    break
        if año_real:
            año_actual = año_real

        # Construir la fecha de semana: preferir columna "Semana" pero corregir el año
        # si difiere del año real de los eventos
        fecha_col_semana = registros[0].get('semana')
        if fecha_col_semana and año_real and fecha_col_semana.year != año_real:
            # El año en la columna Semana es incorrecto — reconstruir con año_real
            fecha_semana = fecha_col_semana.replace(year=año_real)
        elif fecha_col_semana:
            fecha_semana = fecha_col_semana
        else:
            fecha_semana = _fecha_de_hoja(nombre_hoja, año=año_actual)

        if fecha_semana:
            semanas[fecha_semana] = registros

    if not semanas:
        return {'semanas': {}, 'error': 'No se encontraron hojas semanales en el Excel'}

    return {'semanas': semanas, 'error': None}


# ---------------------------------------------------------------------------
# Función principal
# ---------------------------------------------------------------------------

def importar_incidencias(semana: date = None, remitente: str = None) -> dict:
    """
    Busca en Gmail el correo de incidencias (para la semana dada o el más reciente)
    y retorna los datos parseados del Excel adjunto.

    Returns:
        {
          'semanas': { date: [registros] },
          'msg_id': str | None,
          'error': str | None,
        }
    """
    try:
        service = _get_service()
    except Exception as exc:
        return {'semanas': {}, 'msg_id': None, 'error': f'Error de autenticación: {exc}'}

    query = _build_query(semana, remitente=remitente) if semana else _build_query_broad(remitente=remitente)
    logger.info("Gmail query incidencias: %s", query)

    try:
        resultado = service.users().messages().list(
            userId='me', q=query, maxResults=5
        ).execute()
    except Exception as exc:
        return {'semanas': {}, 'msg_id': None, 'error': f'Error al buscar correos: {exc}'}

    mensajes = resultado.get('messages', [])
    if not mensajes:
        return {'semanas': {}, 'msg_id': None, 'error': 'No se encontró correo de incidencias'}

    msg_id = mensajes[0]['id']

    try:
        xlsx_bytes = _get_attachment_xlsx(service, msg_id)
    except Exception as exc:
        return {'semanas': {}, 'msg_id': msg_id, 'error': f'Error al descargar adjunto: {exc}'}

    if not xlsx_bytes:
        return {'semanas': {}, 'msg_id': msg_id, 'error': 'El correo no tiene adjunto .xlsx'}

    resultado_parseo = parsear_excel_incidencias(xlsx_bytes)
    resultado_parseo['msg_id'] = msg_id
    return resultado_parseo
